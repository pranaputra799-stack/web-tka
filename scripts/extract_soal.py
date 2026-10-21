"""
Final robust extractor for all .docx format variants.
Handles different formats:
- Question header styles: "Question X (Single Choice)", "Question X (Pilihan Ganda...)", "Question X" standalone
- Answer key styles: "Kunci Jawaban:", "Jawaban Benar:", "Statement X: (A) Category"
- Text section styles: "LEMBAR SOAL (TEKS X)", "Reading Text", "PART X: TEXT X"
- Option styles: "A. text", "(A) text", "[ ] A. text", inline options, separate paragraphs
"""

import os
import json
import re
import sys
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

BANK_SOAL_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "bank soal")
OUTPUT_FILE = os.path.join(os.path.dirname(os.path.dirname(__file__)), "data", "soal.json")


def get_all_paragraph_texts(doc):
    """Get all paragraph texts with their indices."""
    return [(i, p.text.strip()) for i, p in enumerate(doc.paragraphs)]


def find_question_indices(paragraphs):
    """Find all paragraph indices that start a question."""
    indices = []
    for i, text in enumerate(paragraphs):
        # Match "Question X" at start of line (possibly with type info)
        if re.match(r'^Question\s+\d+', text):
            indices.append(i)
    return indices


def determine_q_type(text_block):
    """Determine question type from surrounding text."""
    t = text_block.lower()
    if any(w in t for w in ['matching table', 'matching', 'klasifikasi tabel', 'kategori',
                             'classify', 'categorize', 'true or false',
                             'cause or an effect', 'cause or effect',
                             'advantage or disadvantage', 'suitable or unsuitable',
                             'agree or disagree', 'argument or explanation',
                             'true or false']):
        return "matching"
    if any(w in t for w in ['more than one', 'every correct answer', 'mcma',
                             'multiple choice', 'pilihan ganda kompleks',
                             'choose every correct', 'click on every correct',
                             'select more than one']):
        return "multiple"
    if any(w in t for w in ['single choice', 'pilihan ganda biasa', 'pilihan ganda tunggal']):
        return "single"
    return "single"


def extract_all_data(filepath):
    """Extract everything from a docx file using a unified approach."""
    doc = Document(filepath)
    all_paras = [p.text for p in doc.paragraphs]
    tables = doc.tables
    
    package_name = all_paras[0].strip() if all_paras else os.path.basename(filepath)
    
    # ========== EXTRACT READING TEXTS ==========
    texts = []
    text_num = 0
    i = 0
    while i < len(all_paras):
        line = all_paras[i].strip()
        
        is_text_header = False
        # Check for various text section markers
        lembar_match = re.match(r'LEMBAR SOAL \(TEKS (\d+)\)', line, re.IGNORECASE)
        part_match = re.match(r'PART \d+:\s*TEXT (\d+)', line, re.IGNORECASE)
        reading_match = re.match(r'Reading Text(?:\s+(\d+))?', line, re.IGNORECASE)
        
        if lembar_match:
            text_num = int(lembar_match.group(1))
            is_text_header = True
        elif part_match:
            text_num = int(part_match.group(1))
            is_text_header = True
        elif reading_match:
            if reading_match.group(1):
                text_num = int(reading_match.group(1))
            else:
                text_num += 1
            is_text_header = True
        
        if is_text_header:
            i += 1
            title = ""
            passage_parts = []
            while i < len(all_paras):
                ptext = all_paras[i].strip()
                if not ptext:
                    i += 1
                    continue
                # Skip metadata lines
                if ptext.startswith("Jenis Teks:") or ptext.startswith("Kompetensi:"):
                    i += 1
                    continue
                # Stop at question or next section
                if (re.match(r'^Question\s+\d+', ptext) or 
                    ptext == "QUESTIONS" or
                    re.match(r'LEMBAR SOAL', ptext) or
                    re.match(r'PART \d+:', ptext) or
                    ptext == "Reading Text"):
                    break
                if not title:
                    title = ptext
                else:
                    passage_parts.append(ptext)
                i += 1
            
            if title:
                texts.append({
                    "text_number": text_num,
                    "title": title,
                    "passage": "\n\n".join(passage_parts)
                })
            continue
        i += 1
    
    # ========== EXTRACT QUESTIONS ==========
    questions = []
    table_idx = 0
    
    # Find all question start positions
    q_positions = []
    for i, line in enumerate(all_paras):
        stripped = line.strip()
        if re.match(r'^Question\s+\d+', stripped):
            # Extract question number
            num_match = re.match(r'^Question\s+(\d+)', stripped)
            if num_match:
                q_positions.append((i, int(num_match.group(1))))
    
    # Also find where answer key sections start
    answer_section_start = len(all_paras)
    for i, line in enumerate(all_paras):
        stripped = line.strip()
        if re.match(r'KUNCI JAWABAN|LEMBAR KUNCI', stripped):
            answer_section_start = i
            break
    
    for q_idx, (para_idx, q_num) in enumerate(q_positions):
        # Skip questions in the answer key section
        if para_idx >= answer_section_start:
            continue
        
        # Determine end of this question (start of next question or answer section)
        end_idx = answer_section_start
        for next_para_idx, _ in q_positions[q_idx + 1:]:
            if next_para_idx > para_idx and next_para_idx < answer_section_start:
                end_idx = next_para_idx
                break
        
        # Gather all text for this question
        q_lines = []
        for j in range(para_idx, min(end_idx, len(all_paras))):
            ptext = all_paras[j].strip()
            if j > para_idx and (re.match(r'LEMBAR SOAL \(TEKS \d+\)', ptext, re.I) or re.match(r'PART \d+:\s*TEXT \d+', ptext, re.I) or re.match(r'Reading Text(?:\s+\d+)?', ptext, re.I)):
                break
            q_lines.append(ptext)
        
        q_block = '\n'.join(q_lines)
        
        # Determine type
        q_type = determine_q_type(q_block)
        
        # Parse question text and options
        q_text = ""
        options = {}
        
        # Split into sublines and filter
        content_lines = []
        for line in q_lines:
            for subline in line.split('\n'):
                subline = subline.strip()
                if not subline:
                    continue
                if re.match(r'^Question\s+\d+', subline):
                    continue
                if subline.startswith('(Indikator') or subline.startswith('Indikator') or subline.startswith('Bentuk Soal:'):
                    continue
                if subline == "QUESTIONS":
                    continue
                content_lines.append(subline)
        
        if q_type == "matching":
            # Question text is the content lines joined
            q_text = ' '.join(content_lines).strip()
            if table_idx < len(tables):
                table = tables[table_idx]
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                # Deduplicate headers (Word tables sometimes merge cells)
                seen = []
                for h in headers:
                    if h not in seen:
                        seen.append(h)
                headers = seen
                
                categories = headers[1:] if len(headers) > 1 else []
                statements = []
                for row in table.rows[1:]:
                    stmt = row.cells[0].text.strip()
                    if stmt:
                        statements.append(stmt)
                options = {
                    "headers": headers,
                    "categories": categories,
                    "statements": statements
                }
                table_idx += 1
        else:
            # Parse question text and options from content_lines
            q_text_parts = []
            
            for line in content_lines:
                # Check for checkbox-style options: [ ] A. text or [x] A. text
                checkbox = re.match(r'\[[ x]?\]\s*([A-E])\.\s+(.+)', line)
                if checkbox:
                    options[checkbox.group(1)] = checkbox.group(2).strip()
                    continue
                
                # Check for parenthesized options: (A) text
                paren_opt = re.match(r'^\(([A-E])\)\s+(.+)', line)
                if paren_opt:
                    options[paren_opt.group(1)] = paren_opt.group(2).strip()
                    continue
                
                # Check for regular options: A. text (at start of line)
                reg_opt = re.match(r'^([A-E])\.\s+(.+)', line)
                if reg_opt:
                    options[reg_opt.group(1)] = reg_opt.group(2).strip()
                    continue
                
                # Try inline options (A. text\nB. text within same paragraph)
                if '\n' in line:
                    sub_lines = line.split('\n')
                    has_opts = False
                    temp_opts = {}
                    temp_text = []
                    for sl in sub_lines:
                        sl = sl.strip()
                        # Check various option formats
                        om = re.match(r'^([A-E])\.\s+(.+)', sl) or re.match(r'^\(([A-E])\)\s+(.+)', sl) or re.match(r'\[[ x]?\]\s*([A-E])\.\s+(.+)', sl)
                        if om:
                            temp_opts[om.group(1)] = om.group(2).strip()
                            has_opts = True
                        elif sl and not any(skip in sl.lower() for skip in ['more than one', 'click on every', 'choose every', 'there is more']):
                            temp_text.append(sl)
                    
                    if has_opts:
                        options.update(temp_opts)
                        q_text_parts.extend(temp_text)
                        continue
                
                # Skip instruction lines
                if any(skip in line.lower() for skip in ['more than one correct answer', 'click on every correct', 'choose every correct answer', 'there is more than one']):
                    continue
                
                q_text_parts.append(line)
            
            q_text = ' '.join(q_text_parts).strip()
        
        text_number = ((q_num - 1) // 5) + 1
        
        questions.append({
            "number": q_num,
            "type": q_type,
            "text_number": text_number,
            "question": q_text,
            "options": options
        })
    
    # ========== EXTRACT ANSWER KEYS ==========
    answer_keys = {}
    
    i = 0
    while i < len(all_paras):
        line = all_paras[i].strip()
        
        # Pattern 1: "X. Kunci Jawaban: A, B, C"
        m1 = re.match(r'^(\d+)\.\s*Kunci Jawaban:\s*([A-E](?:\s*,\s*[A-E])*)', line)
        if m1:
            q_num = int(m1.group(1))
            raw = m1.group(2)
            answers = [a.strip() for a in raw.split(',')]
            
            pembahasan = ""
            p_match = re.search(r'Pembahasan:\s*(.+)', line, re.DOTALL)
            if p_match:
                pembahasan = p_match.group(1).strip()
            elif i + 1 < len(all_paras):
                next_line = all_paras[i + 1].strip()
                if next_line.startswith('Pembahasan:'):
                    pembahasan = next_line[len('Pembahasan:'):].strip()
            
            answer_keys[q_num] = {"answers": answers, "explanation": pembahasan, "type": "choice"}
            i += 1
            continue
        
        # Pattern 1b: "X. Kunci Jawaban: B (Catatan Teknis: ...)"
        m1b = re.match(r'^(\d+)\.\s*Kunci Jawaban:\s*([A-E](?:\s*,\s*[A-E])*)\s*\(', line)
        if m1b:
            q_num = int(m1b.group(1))
            answers = [a.strip() for a in m1b.group(2).split(',')]
            pembahasan = ""
            if i + 1 < len(all_paras):
                next_line = all_paras[i + 1].strip()
                if next_line.startswith('Pembahasan:'):
                    pembahasan = next_line[len('Pembahasan:'):].strip()
            answer_keys[q_num] = {"answers": answers, "explanation": pembahasan, "type": "choice"}
            i += 1
            continue
        
        # Pattern 2: "X. Kunci Jawaban & Klasifikasi Tabel:"
        m2 = re.match(r'^(\d+)\.\s*Kunci Jawaban\s*[&]\s*Klasifikasi Tabel:', line)
        if m2:
            q_num = int(m2.group(1))
            matching = {}
            j = i + 1
            while j < len(all_paras) and j < i + 12:
                stmt = all_paras[j].strip()
                # "Pernyataan 1 Category:" or "Pernyataan 1 -> Category:"
                sm = re.match(r'Pernyataan\s+(\d+)\s+(?:->\s*)?(.+?):', stmt)
                if sm:
                    matching[str(sm.group(1))] = {
                        "category": sm.group(2).strip(),
                        "explanation": stmt[sm.end():].strip()
                    }
                elif stmt and re.match(r'^\d+\.\s*Kunci|LEMBAR|KUNCI', stmt):
                    break
                j += 1
            answer_keys[q_num] = {"answers": matching, "explanation": "", "type": "matching"}
            i = j
            continue
        
        # Pattern 3: "Question X" followed by "Kunci Jawaban:" or "Jawaban Benar:"
        q_ref = re.match(r'^Question\s+(\d+)$', line)
        if q_ref and i >= answer_section_start:
            q_num = int(q_ref.group(1))
            if i + 1 < len(all_paras):
                next_line = all_paras[i + 1].strip()
                
                # "Kunci Jawaban: A, B, C" or "Kunci Jawaban: A, B, C (note)"
                km = re.match(r'Kunci Jawaban:\s*([A-E](?:\s*,\s*[A-E])*)', next_line)
                if km:
                    raw = re.sub(r'\s*\(.*?\)', '', km.group(1)).strip()
                    answers = [a.strip() for a in raw.split(',')]
                    pembahasan = ""
                    if i + 2 < len(all_paras):
                        p_text = all_paras[i + 2].strip()
                        if p_text.startswith('Pembahasan:'):
                            pembahasan = p_text[len('Pembahasan:'):].strip()
                    answer_keys[q_num] = {"answers": answers, "explanation": pembahasan, "type": "choice"}
                    i += 3
                    continue
                
                # "Jawaban Benar: A, B, dan C" or "Jawaban Benar: B"
                jm = re.match(r'Jawaban Benar:\s*(.+)', next_line)
                if jm:
                    raw = jm.group(1).strip()
                    # Parse "A, B, dan C" or "A, C, dan E" format
                    raw = raw.replace(' dan ', ', ').replace(' and ', ', ')
                    letters = re.findall(r'[A-E]', raw)
                    if letters:
                        pembahasan = ""
                        if i + 2 < len(all_paras):
                            p_text = all_paras[i + 2].strip()
                            if p_text.startswith('Pembahasan:'):
                                pembahasan = p_text[len('Pembahasan:'):].strip()
                        answer_keys[q_num] = {"answers": letters, "explanation": pembahasan, "type": "choice"}
                        i += 3
                        continue
                
                # "Kunci Jawaban:" (empty - matching type) or "Jawaban Benar:" (empty)
                if next_line in ("Kunci Jawaban:", "Jawaban Benar:"):
                    matching = {}
                    pembahasan = ""
                    j = i + 2
                    while j < len(all_paras) and j < i + 15:
                        stmt = all_paras[j].strip()
                        # "Statement X: (A) Category"
                        sm1 = re.match(r'Statement\s+(\d+):\s*\(([A-E])\)\s*(.+)', stmt)
                        if sm1:
                            matching[str(sm1.group(1))] = {
                                "category": sm1.group(3).strip(),
                                "letter": sm1.group(2)
                            }
                            j += 1
                            continue
                        
                        # "Pernyataan X -> Category:" 
                        sm2 = re.match(r'Pernyataan\s+(\d+)\s+(?:->\s*)?(.+?):', stmt)
                        if sm2:
                            matching[str(sm2.group(1))] = {
                                "category": sm2.group(2).strip(),
                                "explanation": stmt[sm2.end():].strip()
                            }
                            j += 1
                            continue
                        
                        # "Experiencing diverse... → Advantage"
                        arrow_match = re.search(r'(.+?)\s*(?:\\\\?\(\\\\?rightarrow\\\\?\)|->|→|➔)\s*(.+)', stmt)
                        if arrow_match:
                            stmt_text = arrow_match.group(1).strip()
                            category = arrow_match.group(2).strip()
                            stmt_idx = str(len(matching) + 1)
                            matching[stmt_idx] = {
                                "category": category,
                                "statement_text": stmt_text
                            }
                            j += 1
                            continue
                        
                        if stmt.startswith('Pembahasan:'):
                            pembahasan = stmt[len('Pembahasan:'):].strip()
                            j += 1
                            continue
                        
                        if re.match(r'^Question \d+|^KUNCI', stmt):
                            break
                        j += 1
                    
                    if matching:
                        answer_keys[q_num] = {"answers": matching, "explanation": pembahasan, "type": "matching"}
                    i = j
                    continue
        
        i += 1
    
    # Merge answer keys into questions
    for q in questions:
        if q["number"] in answer_keys:
            ak = answer_keys[q["number"]]
            q["answer_key"] = ak["answers"]
            q["explanation"] = ak.get("explanation", "")
            if ak["type"] == "matching":
                q["matching_answers"] = ak["answers"]
    
    return {
        "package_name": package_name,
        "filename": os.path.basename(filepath),
        "texts": texts,
        "questions": questions,
        "total_questions": len(questions)
    }


def get_sort_key(filename):
    name_lower = filename.lower()
    if 'prediksi' in name_lower:
        numbers = re.findall(r'\d+', filename)
        return (0, int(numbers[0]) if numbers else 0)
    elif 'paket' in name_lower:
        paket_match = re.search(r'paket\s*(\d+)', name_lower)
        if paket_match:
            return (1, int(paket_match.group(1)))
        numbers = re.findall(r'\d+', filename)
        return (1, int(numbers[0]) if numbers else 0)
    else:
        numbers = re.findall(r'\d+', filename)
        return (2, int(numbers[0]) if numbers else 0)


def main():
    docx_files = sorted(
        [f for f in os.listdir(BANK_SOAL_DIR) if f.endswith('.docx')],
        key=get_sort_key
    )
    
    print(f"Found {len(docx_files)} files to process\n")
    
    all_packages = []
    
    for idx, filename in enumerate(docx_files, 1):
        filepath = os.path.join(BANK_SOAL_DIR, filename)
        print(f"[{idx}/{len(docx_files)}] {filename}")
        
        try:
            package = extract_all_data(filepath)
            package["id"] = idx
            all_packages.append(package)
            
            q_count = package['total_questions']
            t_count = len(package['texts'])
            keys = sum(1 for q in package['questions'] if 'answer_key' in q)
            opts = sum(1 for q in package['questions'] if q.get('options'))
            
            status = "OK" if q_count >= 25 and keys >= 25 else "PARTIAL"
            print(f"  [{status}] {q_count} questions, {t_count} texts, {keys} keys, {opts} with options")
            
        except Exception as e:
            import traceback
            print(f"  [ERROR] {e}")
            traceback.print_exc()
    
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump({
            "total_packages": len(all_packages),
            "packages": all_packages
        }, f, ensure_ascii=False, indent=2)
    
    print(f"\n{'='*60}")
    total_q = sum(p['total_questions'] for p in all_packages)
    total_t = sum(len(p['texts']) for p in all_packages)
    total_k = sum(sum(1 for q in p['questions'] if 'answer_key' in q) for p in all_packages)
    print(f"Total: {total_q} questions, {total_t} texts, {total_k} answer keys")
    print(f"Saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
